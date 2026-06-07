"""Excel exporters for GSTR returns and registers (FR-006, FR-009).

Produces in-memory .xlsx bytes (openpyxl) the API streams as a download. Layout mirrors the
GST offline tool's familiar columns so users/CAs recognise it.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.models.invoice import Invoice

_HEADER_FILL = PatternFill("solid", fgColor="1E3A8A")  # trust blue (SRS design system)
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 3, 45)


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def gstr1_xlsx(summary: dict, sales: list[Invoice]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "GSTR-1 Invoices"
    headers = ["GSTIN/UIN of Recipient", "Invoice Number", "Invoice Date", "Place of Supply",
               "Invoice Value", "Taxable Value", "IGST", "CGST", "SGST", "Cess"]
    ws.append(headers)
    for inv in sales:
        ws.append([
            inv.counterparty_gstin or "B2C",
            inv.invoice_no,
            inv.invoice_date.strftime("%d-%m-%Y") if inv.invoice_date else "",
            inv.place_of_supply or "",
            inv.total_value, inv.taxable_value, inv.igst, inv.cgst, inv.sgst, inv.cess,
        ])
    _style_header(ws)
    _autosize(ws)

    sumws = wb.create_sheet("Summary")
    for k, v in summary.items():
        sumws.append([k, v])
    return _save(wb)


def gstr3b_xlsx(summary: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "GSTR-3B Summary"
    ws.append(["Field", "Value"])

    def emit(prefix: str, obj):
        if isinstance(obj, dict):
            for k, v in obj.items():
                emit(f"{prefix}.{k}" if prefix else k, v)
        else:
            ws.append([prefix, obj])

    emit("", summary)
    _style_header(ws)
    _autosize(ws)
    return _save(wb)


def register_xlsx(title: str, invoices: list[Invoice]) -> bytes:
    """Sales or Purchase register export (FR-009)."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(["Date", "Invoice No", "Party", "GSTIN", "Taxable", "CGST", "SGST",
               "IGST", "Cess", "Total", "Status"])
    for inv in invoices:
        ws.append([
            inv.invoice_date.strftime("%d-%m-%Y") if inv.invoice_date else "",
            inv.invoice_no, inv.counterparty_name, inv.counterparty_gstin,
            inv.taxable_value, inv.cgst, inv.sgst, inv.igst, inv.cess,
            inv.total_value, inv.status,
        ])
    _style_header(ws)
    _autosize(ws)
    return _save(wb)
